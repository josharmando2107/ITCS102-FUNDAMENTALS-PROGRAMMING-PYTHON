import tkinter as tk
from tkinter import messagebox
import openpyxl
from datetime import datetime

def submit_data():
    current_year = datetime.now().year
    
    # Initialize a new workbook and select the active worksheet
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Favorites"
    
    # 5. Define Excel columns
    headers = ["ID", "First Name", "Last Name", "Birth Year", "Age"]
    ws.append(headers)
    
    # Retrieve data from GUI, process it, and write to Excel
    for idx, (entry_fn, entry_ln, entry_by) in enumerate(entries):
        first_name = entry_fn.get().strip()
        last_name = entry_ln.get().strip()
        birth_year_str = entry_by.get().strip()
        
        # Validation
        if not first_name or not last_name or not birth_year_str:
            messagebox.showerror("Input Error", f"Please fill out all fields for Person {idx+1}.")
            return
            
        try:
            birth_year = int(birth_year_str)
        except ValueError:
            messagebox.showerror("Input Error", f"Birth year for Person {idx+1} must be a number.")
            return
            
        # 2. Automatically compute the age
        age = current_year - birth_year
        
        # 3. Assign an ID number
        person_id = idx + 1
        
        # 6. Store all user input in the Excel file
        row_data = [person_id, first_name, last_name, birth_year, age]
        ws.append(row_data)
        
    # 4. Automatically create and save the Excel file
    filename = "favorite_people.xlsx"
    wb.save(filename)
    
    messagebox.showinfo("Success", f"Data successfully saved to {filename}!\nCheck your console for the output.")
    root.destroy() # Close the GUI window
    
    # 7. Display all stored data in the console
    display_console_data(filename)

def display_console_data(filename):
    print("\n--- Stored Data from Excel ---")
    
    # Load the saved workbook to prove data was stored correctly
    saved_wb = openpyxl.load_workbook(filename)
    saved_ws = saved_wb.active
    
    # Print headers and rows neatly
    for row in saved_ws.iter_rows(values_only=True):
        # Format the output with fixed widths for clean console display
        print(f"{str(row[0]):<5} | {str(row[1]):<15} | {str(row[2]):<15} | {str(row[3]):<12} | {str(row[4]):<5}")
    print("------------------------------\n")

# --- GUI Setup ---
root = tk.Tk()
root.title("Favorite People Data Entry")
root.geometry("350x450")
root.padx = 20

# 1. Ask the user to enter information for THREE favorite people
tk.Label(root, text="Enter Details for 3 Favorite People", font=("Arial", 12, "bold")).pack(pady=10)

entries = []

for i in range(3):
    frame = tk.LabelFrame(root, text=f"Person {i+1}", padx=10, pady=5)
    frame.pack(fill="x", padx=10, pady=5)
    
    # First Name
    tk.Label(frame, text="First Name:").grid(row=0, column=0, sticky="w")
    fn_entry = tk.Entry(frame)
    fn_entry.grid(row=0, column=1, pady=2)
    
    # Last Name
    tk.Label(frame, text="Last Name:").grid(row=1, column=0, sticky="w")
    ln_entry = tk.Entry(frame)
    ln_entry.grid(row=1, column=1, pady=2)
    
    # Birth Year
    tk.Label(frame, text="Birth Year:").grid(row=2, column=0, sticky="w")
    by_entry = tk.Entry(frame)
    by_entry.grid(row=2, column=1, pady=2)
    
    entries.append((fn_entry, ln_entry, by_entry))

# Submit Button
submit_btn = tk.Button(root, text="Save to Excel", command=submit_data, bg="green", fg="white", font=("Arial", 10, "bold"))
submit_btn.pack(pady=15)

root.mainloop()
